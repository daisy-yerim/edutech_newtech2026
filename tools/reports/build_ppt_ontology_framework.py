from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path("artifacts/presentations/ppt_ontology_framework.pptx")

NAVY = RGBColor(20, 37, 63)
BLUE = RGBColor(38, 99, 235)
CYAN = RGBColor(14, 165, 233)
MINT = RGBColor(16, 185, 129)
AMBER = RGBColor(245, 158, 11)
RED = RGBColor(239, 68, 68)
INK = RGBColor(30, 41, 59)
MUTED = RGBColor(100, 116, 139)
LIGHT = RGBColor(241, 245, 249)
WHITE = RGBColor(255, 255, 255)


def text_box(slide, x, y, w, h, text, size=18, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(0.04)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "맑은 고딕"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def rounded(slide, x, y, w, h, fill, line=None, radius=True):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def label(slide, x, y, w, text, fill=BLUE):
    rounded(slide, x, y, w, 0.34, fill)
    text_box(slide, x, y, w, 0.34, text, 10, WHITE, True, PP_ALIGN.CENTER)


def title(slide, kicker, heading, sub=None):
    text_box(slide, 0.65, 0.35, 2.5, 0.3, kicker.upper(), 10, BLUE, True)
    text_box(slide, 0.65, 0.68, 12.0, 0.58, heading, 25, NAVY, True)
    if sub:
        text_box(slide, 0.66, 1.25, 12.0, 0.38, sub, 11, MUTED)


def arrow(slide, x1, y1, x2, y2, color=BLUE, width=2.0):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    line.line.end_arrowhead = True
    return line


def footer(slide, page):
    text_box(slide, 0.65, 7.15, 11.8, 0.2, "Korean RAG · Executable Ontology Framework", 8, MUTED)
    text_box(slide, 12.25, 7.15, 0.4, 0.2, str(page), 8, MUTED, False, PP_ALIGN.RIGHT)


def add_overview(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "FRAMEWORK 01", "PPT 기반 한국어 평가 생성 × 실행형 온톨로지",
          "자료를 읽는 RAG에서, 근거·급수·언어 자원을 통제하는 생성 시스템으로")

    stages = [
        ("01", "PPT·원천자료", "슬라이드, 표, 도형,\n발표자 노트 추출", BLUE),
        ("02", "의미 구조화", "개념·사실·관계·\n출처 단위 정규화", CYAN),
        ("03", "온톨로지 결합", "급수·텍스트 유형·\n어휘·문법 연결", MINT),
        ("04", "제약 기반 생성", "지문·문항·답안 생성\n근거와 난이도 통제", AMBER),
        ("05", "이중 검증", "AI 품질 평가 +\n규칙·SHACL 판정", RED),
    ]
    x0, gap, w = 0.66, 0.18, 2.32
    for i, (num, name, desc, color) in enumerate(stages):
        x = x0 + i * (w + gap)
        rounded(s, x, 2.05, w, 2.38, WHITE, RGBColor(226, 232, 240))
        label(s, x + 0.18, 2.25, 0.48, num, color)
        text_box(s, x + 0.18, 2.78, w - 0.36, 0.42, name, 16, NAVY, True)
        text_box(s, x + 0.18, 3.28, w - 0.36, 0.72, desc, 11, MUTED)
        if i < 4:
            arrow(s, x + w, 3.25, x + w + gap, 3.25, MUTED, 1.3)

    rounded(s, 0.66, 4.86, 12.0, 1.55, LIGHT, LIGHT)
    text_box(s, 0.93, 5.08, 2.15, 0.35, "ONTOLOGY CONTROL PLANE", 10, BLUE, True)
    text_box(s, 0.93, 5.48, 11.35, 0.45,
             "GROUNDED_IN  ·  TARGETS_GRADE  ·  HAS_REQUIREMENT  ·  SUPPORTS_TEXT_TYPE  ·  ALLOWS_RESOURCE",
             13, NAVY, True, PP_ALIGN.CENTER)
    text_box(s, 0.93, 5.94, 11.35, 0.25,
             "온톨로지가 전 단계의 입력 조건과 검증 결과를 같은 의미 체계로 연결",
             10, MUTED, False, PP_ALIGN.CENTER)
    footer(s, 1)


def add_model(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "FRAMEWORK 02", "핵심 의미 모델: 무엇을 어떤 관계로 연결하는가",
          "PPT의 표현 단위를 출처가 추적되는 학습·평가 지식으로 변환")

    nodes = {
        "ppt": (0.75, 2.45, 2.1, 0.82, "PPT / SourceDocument", BLUE),
        "fact": (3.55, 1.88, 2.1, 0.82, "Fact · Concept", CYAN),
        "passage": (3.55, 3.30, 2.1, 0.82, "Passage", MINT),
        "grade": (6.70, 1.55, 2.1, 0.82, "Grade", AMBER),
        "resource": (9.95, 1.55, 2.35, 0.82, "Lexeme · Grammar", CYAN),
        "type": (6.70, 3.18, 2.1, 0.82, "TextType", AMBER),
        "item": (9.95, 3.18, 2.35, 0.82, "Question · Answer", RED),
        "evidence": (6.70, 4.82, 2.1, 0.82, "Evidence", BLUE),
        "validation": (9.95, 4.82, 2.35, 0.82, "ValidationResult", RED),
    }
    for x, y, w, h, txt, color in nodes.values():
        rounded(s, x, y, w, h, WHITE, color)
        text_box(s, x + 0.08, y, w - 0.16, h, txt, 13, color, True, PP_ALIGN.CENTER)

    relations = [
        (2.85, 2.73, 3.55, 2.30, "CONTAINS_FACT"),
        (2.85, 2.99, 3.55, 3.68, "GROUNDS"),
        (5.65, 3.52, 6.70, 1.96, "TARGETS_GRADE"),
        (5.65, 3.72, 6.70, 3.59, "REALIZES"),
        (8.80, 1.96, 9.95, 1.96, "ALLOWS_RESOURCE"),
        (8.80, 3.59, 9.95, 3.59, "HAS_ITEM"),
        (5.65, 3.95, 6.70, 5.22, "SUPPORTED_BY"),
        (8.80, 5.22, 9.95, 5.22, "PRODUCES"),
    ]
    for x1, y1, x2, y2, rel in relations:
        arrow(s, x1, y1, x2, y2, MUTED, 1.2)
        mx, my = (x1 + x2) / 2, (y1 + y2) / 2
        text_box(s, mx - 0.55, my - 0.18, 1.1, 0.28, rel, 7, MUTED, True, PP_ALIGN.CENTER)

    text_box(s, 0.78, 5.93, 11.7, 0.44,
             "설계 원칙  |  출처 추적성 · 급수 적합성 · 자원 허용성 · 문항 정합성 · 판정 재현성",
             12, NAVY, True, PP_ALIGN.CENTER)
    footer(s, 2)


def add_loop(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "FRAMEWORK 03", "생성–검증–수정의 폐쇄 루프",
          "규칙 위반을 점수로 끝내지 않고 다음 생성 라운드의 구체적 수정 조건으로 환류")

    cols = [
        (0.75, "1. RETRIEVE", "근거 선택", "PPT 사실·개념\n급수별 허용 자원", BLUE),
        (3.25, "2. GENERATE", "초안 생성", "지문·문항·답안\n출처 ID 포함", MINT),
        (5.75, "3. VALIDATE", "이중 검증", "AI: 사실·자연스러움\nOntology: 길이·급수·관계", AMBER),
        (8.25, "4. REPAIR", "원인 기반 수정", "위반 어휘 교체\n길이·근거·문항 보정", RED),
        (10.75, "5. APPROVE", "승인·배포", "교사 검토 기록\n버전·판정 이력 보존", BLUE),
    ]
    for i, (x, tag, name, desc, color) in enumerate(cols):
        rounded(s, x, 2.17, 1.88, 2.62, WHITE, RGBColor(226, 232, 240))
        label(s, x + 0.15, 2.38, 1.58, tag, color)
        text_box(s, x + 0.15, 2.92, 1.58, 0.42, name, 15, NAVY, True, PP_ALIGN.CENTER)
        text_box(s, x + 0.15, 3.49, 1.58, 0.76, desc, 10, MUTED, False, PP_ALIGN.CENTER)
        if i < 4:
            arrow(s, x + 1.88, 3.48, x + 2.50, 3.48, MUTED, 1.4)

    rounded(s, 3.25, 5.22, 7.38, 0.84, LIGHT, LIGHT)
    text_box(s, 3.47, 5.30, 6.94, 0.28,
             "FAIL → 위반 코드 + 대상 노드 + 수정 힌트", 12, RED, True, PP_ALIGN.CENTER)
    text_box(s, 3.47, 5.61, 6.94, 0.25,
             "O-VAL-LENGTH · O-VAL-GRADE-RESOURCE · P-VAL-FACTUALITY", 9, MUTED, False, PP_ALIGN.CENTER)
    arrow(s, 8.90, 5.20, 4.10, 4.82, RED, 1.7)

    text_box(s, 0.82, 6.42, 11.65, 0.31,
             "통과 조건 = AI 품질 검증 PASS ∩ 온톨로지 제약 PASS ∩ 교사 승인",
             13, NAVY, True, PP_ALIGN.CENTER)
    footer(s, 3)


def add_roadmap(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "FRAMEWORK 04", "구현 범위와 성과 지표",
          "현재 자산을 재사용하면서 추적성과 자동 검증을 단계적으로 강화")

    phases = [
        ("NOW", "기본 연결", ["PPT/원문–지문 연결", "급수·길이·텍스트 유형", "어휘 자원 검증"], BLUE),
        ("NEXT", "근거 정밀화", ["생성 문장–원문 사실 대응", "상위어–쉬운 대체어", "문항–근거 문장 연결"], MINT),
        ("SCALE", "운영 자동화", ["SHACL 정책 배포", "검증 이력 대시보드", "전문가 수정의 정책 반영"], AMBER),
    ]
    for i, (tag, name, bullets, color) in enumerate(phases):
        x = 0.75 + i * 4.18
        rounded(s, x, 2.05, 3.72, 2.67, WHITE, RGBColor(226, 232, 240))
        label(s, x + 0.20, 2.28, 0.78, tag, color)
        text_box(s, x + 0.20, 2.82, 3.30, 0.38, name, 16, NAVY, True)
        for j, bullet in enumerate(bullets):
            text_box(s, x + 0.24, 3.30 + j * 0.42, 3.20, 0.30, f"• {bullet}", 10, MUTED)

    metrics = [
        ("Traceability", "근거 연결률", "생성 문장 중 출처 사실과 연결된 비율"),
        ("Conformance", "제약 통과율", "급수·길이·자원·관계 규칙 PASS 비율"),
        ("Efficiency", "수정 효율", "최종 승인까지 평균 재생성 라운드"),
        ("Reliability", "판정 일치도", "AI·규칙·전문가 판정의 일치 수준"),
    ]
    for i, (en, ko, desc) in enumerate(metrics):
        x = 0.75 + i * 3.13
        rounded(s, x, 5.14, 2.78, 1.03, LIGHT, LIGHT)
        text_box(s, x + 0.14, 5.25, 2.50, 0.25, en, 9, BLUE, True)
        text_box(s, x + 0.14, 5.50, 2.50, 0.28, ko, 12, NAVY, True)
        text_box(s, x + 0.14, 5.80, 2.50, 0.22, desc, 7, MUTED)
    footer(s, 4)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "PPT 기반 한국어 평가 생성 × 실행형 온톨로지 프레임워크"
    prs.core_properties.subject = "PPT ingestion, ontology control, constrained generation, validation loop"
    prs.core_properties.author = "Korean RAG Project"
    add_overview(prs)
    add_model(prs)
    add_loop(prs)
    add_roadmap(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
